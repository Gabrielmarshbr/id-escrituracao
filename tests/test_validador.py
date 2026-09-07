from docx import Document

from aportes.documentos.validador import tags_da_minuta, validar_minuta


def _minuta(tmp_path, *paragrafos, nome="m.docx"):
    doc = Document()
    for p in paragrafos:
        doc.add_paragraph(p)
    caminho = tmp_path / nome
    doc.save(caminho)
    return caminho


def test_encontra_as_tags(tmp_path):
    c = _minuta(tmp_path, "Subscritor: {{ nome_subscritor }}", "CPF: {{ cpf_cnpj }}")
    assert tags_da_minuta(c) == {"nome_subscritor", "cpf_cnpj"}


def test_minuta_correta_nao_tem_problema(tmp_path):
    c = _minuta(tmp_path, "{% if pessoa_juridica %}Rep: {{ rep_nome }}{% endif %}")
    assert validar_minuta(c) == []


def test_denuncia_aspas_curvas(tmp_path):
    """A autocorrecao do Word troca as aspas e quebra a tag em silencio."""
    c = _minuta(tmp_path, "{% if pessoa_juridica == “sim” %}x{% endif %}")
    problemas = validar_minuta(c)
    assert any("curva" in p.lower() for p in problemas)


def test_denuncia_condicional_sem_fechamento(tmp_path):
    c = _minuta(tmp_path, "{% if pessoa_juridica %}Rep: {{ rep_nome }}")
    problemas = validar_minuta(c)
    assert any("endif" in p.lower() for p in problemas)


def test_denuncia_chave_desbalanceada(tmp_path):
    c = _minuta(tmp_path, "Subscritor: {{ nome_subscritor }")
    problemas = validar_minuta(c)
    assert any("chave" in p.lower() for p in problemas)


def test_le_tags_dentro_de_tabela(tmp_path):
    """O quadro de integralizacao do Boletim e uma tabela."""
    doc = Document()
    tabela = doc.add_table(rows=1, cols=1)
    tabela.cell(0, 0).text = "Preco de Subscricao: R$ {{ valor_reais }}"
    caminho = tmp_path / "t.docx"
    doc.save(caminho)
    assert "valor_reais" in tags_da_minuta(caminho)
