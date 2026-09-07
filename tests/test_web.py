"""A porta B, ponta a ponta: formulario -> regras -> documentos.

Testa o que a pessoa ve, nao o HTML: que boleta liberada gera PDF, que boleta
travada nao gera nada, e que a tela nunca escreve regra.
"""

from datetime import date

import pytest
from docx import Document

from aportes.configuracao import Configuracao
from aportes.regras import BOLETIM, TERMO_ADESAO
from aportes.web import criar_aplicacao

FUNDOS = """
- id: F1/A
  fundo_nome: Fundo Exemplo FIDC
  fundo_cnpj: "11.111.111/0001-11"
  nome: Classe A
  condominio: fechado
  qualificacao_exigida: qualificado
  fonte: regulamento
  lido_em: 2026-01-15
"""

OFERTAS_PUBLICA = """
- id: OF1
  classe_id: F1/A
  publica: true
"""

OFERTAS_PRIVADA = """
- id: OF1
  classe_id: F1/A
  publica: false
"""


class ConversorFalso:
    def converter(self, docx, pdf):
        pdf.write_bytes(b"%PDF-falso")


@pytest.fixture
def cliente(tmp_path):
    (tmp_path / "regras").mkdir()
    (tmp_path / "regras" / "fundos.yaml").write_text(FUNDOS, encoding="utf-8")
    (tmp_path / "regras" / "ofertas.yaml").write_text(OFERTAS_PUBLICA, encoding="utf-8")
    modelos = tmp_path / "modelos"
    modelos.mkdir()
    for nome in (BOLETIM, TERMO_ADESAO):
        doc = Document()
        doc.add_paragraph("Subscritor: {{ nome_subscritor }}")
        doc.add_paragraph("Valor: R$ {{ valor_reais }} ({{ valor_extenso }})")
        doc.save(modelos / f"{nome}.docx")

    config = Configuracao(raiz=tmp_path, pasta_dados=tmp_path / "dados",
                          pasta_saida=tmp_path / "saida")
    app = criar_aplicacao(config, usuario="maria", conversor=ConversorFalso())
    app.config["TESTING"] = True
    cliente = app.test_client()
    cliente.raiz = tmp_path
    return cliente


def _cadastrar(cliente, categoria="profissional", documento="11111111111"):
    return cliente.post("/cotista", data={
        "documento": documento, "nome": "Fulano de Tal", "tipo": "pf",
        "categoria": categoria, "endereco": "Rua Ficticia, 1",
        "email": "fulano@exemplo.invalido", "fonte": "ficha do Portal ID",
    }, follow_redirects=True)


def _boleta(cliente, **extra):
    dados = {"documento_cotista": "111.111.111-11", "classe_id": "F1/A",
             "oferta_id": "OF1", "valor": "10.000,00", "id_boleta": "B1",
             "data": "2026-09-07", "primeiro_aporte": "sim"}
    dados.update(extra)
    return cliente.post("/avaliar", data=dados)


def test_a_tela_inicial_abre(cliente):
    resposta = cliente.get("/")
    assert resposta.status_code == 200
    assert "Fundo Exemplo FIDC" in resposta.text


def test_cotista_desconhecido_vira_pendencia_com_link(cliente):
    texto = _boleta(cliente).text
    assert "Faltou dado" in texto
    assert "Cadastrar este cotista" in texto


def test_cadastro_e_depois_boleta_gera_os_pdfs(cliente):
    _cadastrar(cliente)
    resposta = _boleta(cliente)

    assert "Liberada" in resposta.text
    pdfs = sorted(p.name for p in (cliente.raiz / "saida").rglob("*.pdf"))
    assert len(pdfs) == 2
    assert any(BOLETIM in nome for nome in pdfs)
    assert any(TERMO_ADESAO in nome for nome in pdfs)


def test_aporte_subsequente_nao_gera_termo_de_adesao(cliente):
    _cadastrar(cliente)
    _boleta(cliente, primeiro_aporte="nao")
    pdfs = [p.name for p in (cliente.raiz / "saida").rglob("*.pdf")]
    assert len(pdfs) == 1
    assert BOLETIM in pdfs[0]


def test_qualificacao_insuficiente_nao_gera_nada(cliente):
    _cadastrar(cliente, categoria="geral")
    resposta = _boleta(cliente)
    assert "Não elegível" in resposta.text
    assert list((cliente.raiz / "saida").rglob("*.pdf")) == []


def test_oferta_privada_para_e_nao_gera_nada(cliente):
    (cliente.raiz / "regras" / "ofertas.yaml").write_text(
        OFERTAS_PRIVADA, encoding="utf-8")
    _cadastrar(cliente)
    resposta = _boleta(cliente)
    assert "análise humana" in resposta.text
    assert list((cliente.raiz / "saida").rglob("*.pdf")) == []


def test_sem_saldo_e_sem_resposta_nao_adivinha(cliente):
    _cadastrar(cliente)
    resposta = _boleta(cliente, primeiro_aporte="")
    assert "Faltou dado" in resposta.text
    assert list((cliente.raiz / "saida").rglob("*.pdf")) == []


def test_mesma_boleta_duas_vezes_nao_duplica(cliente):
    _cadastrar(cliente)
    _boleta(cliente)
    resposta = _boleta(cliente)
    assert "já foi processada" in resposta.text
    assert len(list((cliente.raiz / "saida").rglob("*.pdf"))) == 2


def test_valor_invalido_explica_em_vez_de_quebrar(cliente):
    resposta = _boleta(cliente, valor="dez mil")
    assert resposta.status_code == 400
    assert "valor invalido" in resposta.text


def test_a_tela_nunca_escreve_regra(cliente):
    """Quem decide que um fundo e fechado e quem responde pela area."""
    antes = (cliente.raiz / "regras" / "fundos.yaml").read_text(encoding="utf-8")
    _cadastrar(cliente)
    _boleta(cliente)
    assert (cliente.raiz / "regras" / "fundos.yaml").read_text(
        encoding="utf-8") == antes


def test_cada_pessoa_grava_no_proprio_arquivo(cliente):
    _cadastrar(cliente)
    assert (cliente.raiz / "dados" / "cotistas.maria.json").exists()


def test_documento_invalido_no_cadastro_e_recusado(cliente):
    resposta = cliente.post("/cotista", data={
        "documento": "123", "nome": "X", "tipo": "pf", "categoria": "geral"})
    assert resposta.status_code == 400
    assert "11 digitos" in resposta.text
