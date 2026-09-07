from datetime import date

import pytest
from docx import Document

from aportes.documentos.emissao import (
    MinutaAusente, dados_para_minuta, emitir, nome_base,
)
from aportes.documentos.preenchimento import CampoFaltando
from aportes.dominio import Cotista, Procedencia, TipoPessoa
from aportes.qualificacao import Categoria
from aportes.regras import BOLETIM, TERMO_ADESAO
from tests.conftest import faz_boleta, faz_classe, faz_oferta

PROC = Procedencia(fonte="ficha_pdf", em=date(2026, 3, 12))

CADASTRO_PF = {
    "endereco": "Rua Ficticia, 1 - Sao Paulo/SP",
    "email": "fulano@exemplo.invalido",
    "banco": "001", "agencia": "1234", "conta": "56789-0",
}


def _cotista(tipo=TipoPessoa.PF, cadastro=None, documento="11111111111"):
    return Cotista(
        documento=documento, nome="Fulano de Tal", tipo=tipo,
        categoria=Categoria.PROFISSIONAL, categoria_procedencia=PROC,
        cadastro=dict(cadastro if cadastro is not None else CADASTRO_PF),
        cadastro_procedencia=PROC,
    )


# --- os campos que a minuta pode usar -------------------------------------

def test_campos_de_pessoa_fisica():
    d = dados_para_minuta(faz_boleta(), _cotista(), faz_classe(), faz_oferta())
    assert d["nome_subscritor"] == "Fulano de Tal"
    assert d["cpf_cnpj"] == "111.111.111-11"
    assert d["pessoa_juridica"] == ""
    assert d["endereco"] == "Rua Ficticia, 1 - Sao Paulo/SP"


def test_campos_de_pessoa_juridica():
    cadastro = dict(CADASTRO_PF, representante_nome="Sicrana de Tal",
                    representante_cpf="22222222222")
    d = dados_para_minuta(
        faz_boleta(documento="11111111000111"),
        _cotista(TipoPessoa.PJ, cadastro, documento="11111111000111"),
        faz_classe(), faz_oferta(),
    )
    assert d["cpf_cnpj"] == "11.111.111/0001-11"
    assert d["pessoa_juridica"] == "sim"
    assert d["representante_nome"] == "Sicrana de Tal"
    assert d["representante_cpf"] == "222.222.222-22"


def test_valor_numerico_e_por_extenso():
    d = dados_para_minuta(
        faz_boleta(valor="1234.56"), _cotista(), faz_classe(), faz_oferta())
    assert d["valor_reais"] == "1.234,56"
    assert "mil duzentos e trinta e quatro reais" in d["valor_extenso"]


def test_campos_do_fundo_e_da_classe():
    d = dados_para_minuta(faz_boleta(), _cotista(), faz_classe(), faz_oferta())
    assert d["fundo_nome"] == "Fundo Exemplo FIDC"
    assert d["fundo_cnpj"] == "11.111.111/0001-11"
    assert d["classe_nome"] == "Classe A"
    assert d["data_boleta"] == "07/09/2026"


def test_campo_ausente_no_cadastro_vem_vazio():
    """Vazio para o preenchimento recusar, em vez de inventar."""
    d = dados_para_minuta(
        faz_boleta(), _cotista(cadastro={}), faz_classe(), faz_oferta())
    assert d["endereco"] == ""


# --- nome do arquivo ------------------------------------------------------

def test_prefixo_do_arquivo_e_previsivel():
    prefixo = nome_base(faz_boleta(), _cotista(), faz_classe())
    assert prefixo.startswith("2026-09-07")
    assert "fulano-de-tal" in prefixo
    assert "f1-a" in prefixo


# --- emissao --------------------------------------------------------------

class ConversorFalso:
    def __init__(self):
        self.convertidos = []

    def converter(self, docx, pdf):
        pdf.write_bytes(b"%PDF-falso")
        self.convertidos.append((docx, pdf))


def _minutas(tmp_path, texto="Subscritor: {{ nome_subscritor }}"):
    pasta = tmp_path / "modelos"
    pasta.mkdir()
    for nome in (f"{BOLETIM}.docx", f"{TERMO_ADESAO}.docx"):
        doc = Document()
        doc.add_paragraph(texto)
        doc.save(pasta / nome)
    return pasta


def test_emite_os_documentos_do_veredito(tmp_path):
    modelos = _minutas(tmp_path)
    saida = tmp_path / "saida"
    conversor = ConversorFalso()

    gerados = emitir(
        documentos=(BOLETIM, TERMO_ADESAO),
        dados={"nome_subscritor": "Fulano de Tal"},
        modelos=modelos, saida=saida, prefixo="2026-09-07-fulano",
        conversor=conversor,
    )

    assert len(gerados) == 2
    assert all(p.suffix == ".pdf" and p.exists() for p in gerados)


def test_minuta_ausente_falha_alto(tmp_path):
    modelos = tmp_path / "modelos"
    modelos.mkdir()
    with pytest.raises(MinutaAusente) as erro:
        emitir(documentos=(BOLETIM,), dados={}, modelos=modelos,
               saida=tmp_path / "saida", prefixo="x",
               conversor=ConversorFalso())
    assert BOLETIM in str(erro.value)


def test_campo_faltando_nao_deixa_nada_pela_metade(tmp_path):
    """Se o segundo documento falta dado, o primeiro nao pode ficar solto."""
    modelos = _minutas(tmp_path)
    doc = Document()
    doc.add_paragraph("{{ nome_subscritor }} - {{ campo_que_falta }}")
    doc.save(modelos / f"{TERMO_ADESAO}.docx")
    saida = tmp_path / "saida"

    with pytest.raises(CampoFaltando):
        emitir(documentos=(BOLETIM, TERMO_ADESAO),
               dados={"nome_subscritor": "Fulano de Tal"},
               modelos=modelos, saida=saida, prefixo="x",
               conversor=ConversorFalso())

    assert list(saida.glob("*.pdf")) == []


def test_nao_sobrescreve_documento_existente(tmp_path):
    modelos = _minutas(tmp_path)
    saida = tmp_path / "saida"
    dados = {"nome_subscritor": "Fulano de Tal"}
    emitir(documentos=(BOLETIM,), dados=dados, modelos=modelos, saida=saida,
           prefixo="x", conversor=ConversorFalso())

    with pytest.raises(FileExistsError):
        emitir(documentos=(BOLETIM,), dados=dados, modelos=modelos,
               saida=saida, prefixo="x", conversor=ConversorFalso())
