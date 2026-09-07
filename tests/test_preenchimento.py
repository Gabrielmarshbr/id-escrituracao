import pytest
from docx import Document

from aportes.documentos.preenchimento import CampoFaltando, preencher


def _minuta(tmp_path, *paragrafos):
    doc = Document()
    for p in paragrafos:
        doc.add_paragraph(p)
    caminho = tmp_path / "minuta.docx"
    doc.save(caminho)
    return caminho


def _texto(caminho):
    return "\n".join(p.text for p in Document(caminho).paragraphs)


def test_preenche_as_tags(tmp_path):
    minuta = _minuta(tmp_path, "Subscritor: {{ nome_subscritor }}",
                     "Valor: R$ {{ valor_reais }}")
    destino = tmp_path / "saida.docx"
    preencher(minuta, {"nome_subscritor": "Fulano de Tal",
                       "valor_reais": "10.000,00"}, destino)
    texto = _texto(destino)
    assert "Fulano de Tal" in texto
    assert "10.000,00" in texto


def test_nenhuma_tag_sobra_no_documento(tmp_path):
    minuta = _minuta(tmp_path, "Subscritor: {{ nome_subscritor }}")
    destino = tmp_path / "saida.docx"
    preencher(minuta, {"nome_subscritor": "Fulano de Tal"}, destino)
    texto = _texto(destino)
    assert "{{" not in texto and "}}" not in texto


def test_campo_faltando_impede_a_geracao(tmp_path):
    minuta = _minuta(tmp_path, "{{ nome_subscritor }} - {{ endereco }}")
    destino = tmp_path / "saida.docx"
    with pytest.raises(CampoFaltando) as erro:
        preencher(minuta, {"nome_subscritor": "Fulano de Tal"}, destino)
    assert "endereco" in str(erro.value)
    assert not destino.exists(), "documento incompleto nao pode ser gravado"


def test_campo_vazio_conta_como_faltando(tmp_path):
    """String vazia no Boletim e um buraco, nao um valor."""
    minuta = _minuta(tmp_path, "{{ endereco }}")
    destino = tmp_path / "saida.docx"
    with pytest.raises(CampoFaltando):
        preencher(minuta, {"endereco": "  "}, destino)


def test_condicional_do_word_funciona(tmp_path):
    minuta = _minuta(
        tmp_path,
        "{% if pessoa_juridica %}Representante: {{ rep_nome }}{% endif %}",
    )
    destino = tmp_path / "saida.docx"
    preencher(minuta, {"pessoa_juridica": "", "rep_nome": "-"}, destino)
    assert "Representante" not in _texto(destino)


def test_dado_extra_nao_atrapalha(tmp_path):
    minuta = _minuta(tmp_path, "{{ nome_subscritor }}")
    destino = tmp_path / "saida.docx"
    preencher(minuta, {"nome_subscritor": "Fulano de Tal", "sobra": "x"}, destino)
    assert "Fulano de Tal" in _texto(destino)
