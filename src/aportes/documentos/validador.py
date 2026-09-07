"""Confere uma minuta .docx antes de ela ser usada num bloco.

Existe por causa da autocorrecao do Word, que quebra tags em silencio.
"""

import re
from pathlib import Path

from docx import Document

_TAG = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}")
_ASPAS_CURVAS = "“”‘’"


def _texto_todo(caminho: Path) -> str:
    doc = Document(caminho)
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                partes.append(celula.text)
    return "\n".join(partes)


def tags_da_minuta(caminho: Path) -> set[str]:
    return set(_TAG.findall(_texto_todo(caminho)))


def validar_minuta(caminho: Path) -> list[str]:
    texto = _texto_todo(caminho)
    problemas: list[str] = []

    curvas = [c for c in _ASPAS_CURVAS if c in texto]
    if curvas:
        problemas.append(
            f"aspa curva encontrada ({''.join(curvas)}): a autocorrecao do Word "
            "trocou a aspa reta e isso quebra a tag. Desative a autocorrecao "
            "de aspas e corrija."
        )

    if texto.count("{{") != texto.count("}}"):
        problemas.append(
            f"chave desbalanceada: {texto.count('{{')} aberturas '{{{{' para "
            f"{texto.count('}}')} fechamentos '}}}}'"
        )

    abre = len(re.findall(r"\{%\s*if\b", texto))
    fecha = len(re.findall(r"\{%\s*endif\s*%\}", texto))
    if abre != fecha:
        problemas.append(
            f"condicional sem fechamento: {abre} '{{% if %}}' para "
            f"{fecha} '{{% endif %}}'"
        )

    abre_for = len(re.findall(r"\{%\s*for\b", texto))
    fecha_for = len(re.findall(r"\{%\s*endfor\s*%\}", texto))
    if abre_for != fecha_for:
        problemas.append(
            f"repeticao sem fechamento: {abre_for} '{{% for %}}' para "
            f"{fecha_for} '{{% endfor %}}'"
        )

    return problemas


_TOKEN = re.compile(
    r"\{\{\s*(?P<tag>[a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}"
    r"|\{%\s*if\s+(?P<condicao>.+?)\s*%\}"
    r"|\{%\s*(?P<fim>endif)\s*%\}"
)
_VARIAVEL_SIMPLES = re.compile(r"^[a-zA-Z_][a-zA-Z0-9_]*$")


def tags_exigidas(caminho: Path, dados: dict[str, str]) -> set[str]:
    """As tags que realmente vao aparecer no documento, dados estes valores.

    Uma tag dentro de `{% if x %}` com `x` vazio nao sai no documento, e por
    isso nao pode ser exigida: exigi-la reprovaria, por exemplo, toda pessoa
    fisica por "falta representante legal".

    Condicao que nao seja uma variavel simples e tratada como ativa: na duvida,
    exige o campo. Melhor parar do que gerar documento com buraco.
    """
    exigidas: set[str] = set()
    pilha: list[bool] = []
    for achado in _TOKEN.finditer(_texto_todo(caminho)):
        if achado.group("fim"):
            if pilha:
                pilha.pop()
        elif achado.group("condicao") is not None:
            pilha.append(_condicao_ativa(achado.group("condicao"), dados))
        elif all(pilha):
            exigidas.add(achado.group("tag"))
    return exigidas


def _condicao_ativa(condicao: str, dados: dict[str, str]) -> bool:
    condicao = condicao.strip()
    if _VARIAVEL_SIMPLES.match(condicao):
        return bool(str(dados.get(condicao, "")).strip())
    return True
