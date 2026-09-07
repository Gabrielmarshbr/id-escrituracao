"""Teste de fumaca do Word COM. Rodar A MAO, e OLHAR o PDF gerado.

Nao e automatizavel de forma confiavel: depende do Word instalado.
Na maquina da empresa, este e o primeiro script a rodar.

    .venv/Scripts/python.exe scripts/fumaca_pdf.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document  # noqa: E402

from aportes.documentos.pdf import ConversorWord  # noqa: E402

SAIDA = Path(__file__).resolve().parent.parent / "saida" / "fumaca"


def main() -> None:
    SAIDA.mkdir(parents=True, exist_ok=True)
    docx = SAIDA / "fumaca.docx"
    pdf = SAIDA / "fumaca.pdf"

    doc = Document()
    doc.add_heading("Teste de fumaca", level=1)
    doc.add_paragraph("Se voce esta lendo isto em PDF, o Word COM funciona.")
    doc.add_paragraph("Acentuacao: cotas, subscricao, integralizacao, R$ 1.234,56")
    doc.save(docx)

    with ConversorWord() as conversor:
        conversor.converter(docx, pdf)

    print(f"PDF gerado em: {pdf}")
    print("ABRA O ARQUIVO E CONFIRA. Este teste so vale olhado por gente.")


if __name__ == "__main__":
    main()
